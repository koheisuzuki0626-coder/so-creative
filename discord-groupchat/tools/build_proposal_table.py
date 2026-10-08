#!/usr/bin/env python3
"""insights/ の全リサーチレポートから、提案資料用の一覧表を組む。

各動画について【事実】（題名・投稿チャンネル・再生数・尺・URL・視聴判定）と
【推定】（発注企業・制作会社・動画の目的）を分けて出す。

⚠️ 目的はレポート本文のどこにも書かれていない。検索お題（＝動画の種類）と
題名の語から機械的に導いているだけなので、列名に「推定」と入れ、当たった語を
「推定の根拠」に併記してある。**これを「分析結果」として配らないこと。**

参考外の判定（ボイスサンプル・BGM見本・ポートフォリオ・解説教材・ショート等）は
ai_group_chat.py の規則をそのまま import して使う。写しを作ると、ボット側を
直したときにこちらが古いままになる。

  ./venv/bin/python tools/build_proposal_table.py --dry   # 要約だけ
  ./venv/bin/python tools/build_proposal_table.py         # xlsx と csv を書く
"""
import argparse
import csv
import pathlib
import re
import statistics
import sys
from collections import Counter, defaultdict

BASE = pathlib.Path(__file__).resolve().parent.parent
INSIGHTS = BASE / "insights"
OUT_DIR = BASE.parent / "成果物" / "YouTubeリサーチ提案資料"
sys.path.insert(0, str(BASE))

import ai_group_chat as bot  # noqa: E402  除外規則の唯一の出どころ

# 見出し：## 題名（チャンネル / 12,345回 / 67秒）
# チャンネル名に括弧が入れ子で出るので .+ の貪欲で最後の「/ N回」まで送る。
HEAD_RE = re.compile(
    r"^## (?P<title>.+)（(?P<ch>.+) / (?P<views>[\d,]+)回"
    r"(?: / (?P<dur>\d+)秒)?）\s*$")
URL_RE = re.compile(r"https://www\.youtube\.com/watch\?v=([\w-]+)")
META_RE = re.compile(r"<!-- video id=(?P<id>\S+)(?: duration=(?P<dur>\d+))?"
                     r"(?: verdict=(?P<verdict>\S+))? -->")
TOPIC_RE = re.compile(r"^# YouTube[「【]?(?P<t>.*?)[」】]?\s*リサーチ", re.M)
DATE_RE = re.compile(r"(20\d\d-\d\d-\d\d)")

# ---------- 動画の種類（46通りある検索お題を12種に束ねる） ----------
MISC = "参考外（急上昇・題材の取り違え）"
GENRE_RULES = [
    # 題材の抜き出しが壊れた日のレポート（「今」「試し」「実行」で検索された）。
    # 中身は無関係な動画なので、ジャンルを当てずに参考外へ送る。
    (MISC, r"^(今|試し|実行|他の動画ない|YouTube|YouTubeの動画を|"
           r"YouTubeリサーチもう一回やって)$"),
    (MISC, r"急上昇|トップ100|AI動画|最新のai|AI動画生成"),
    ("MV・ミュージックビデオ", r"ミュージックビデオ|MV|PV|illit|ブラックピンク|椎名林檎|名作"),
    ("WebCM", r"WebCM|webcm"),
    ("採用動画", r"採用"),
    ("会社紹介・コーポレート", r"会社紹介|コーポレート|企業VP|企業プロモーション|"
                              r"会社プロモーション|ブランディング|会社 ?密着"),
    ("サービス・商品紹介", r"サービス紹介|商品紹介|製品紹介"),
    ("アニメーション説明動画", r"アニメーション|説明動画"),
    ("SNS広告（縦型）", r"SNS広告|縦型"),
    ("展示会・サイネージ", r"展示会|サイネージ"),
    ("マニュアル動画", r"マニュアル"),
    ("研修動画", r"研修"),
    ("社内報動画", r"社内報"),
]

# ---------- 目的（推定） ----------
GENRE_PURPOSE = {
    "MV・ミュージックビデオ": "楽曲のプロモーション（音楽作品の発表）",
    "WebCM": "商品・サービスの認知獲得（広告出稿）",
    "採用動画": "採用（応募者の集客・入社前の理解）",
    "会社紹介・コーポレート": "会社の信頼獲得（営業・採用の共通素材）",
    "サービス・商品紹介": "サービス・商品の理解促進（検討中の見込み客向け）",
    "アニメーション説明動画": "仕組みの説明（文字では伝わらない内容の可視化）",
    "SNS広告（縦型）": "SNSでの獲得（短尺・縦型の広告）",
    "展示会・サイネージ": "展示会ブースの集客（通行者の足止め）",
    "マニュアル動画": "操作・業務手順の教育（問い合わせ削減）",
    "研修動画": "社内教育（研修の標準化）",
    "社内報動画": "社内への情報共有（理念・活動の浸透）",
    MISC: "不明（企業映像ではない）",
}
# 上から順に当たった語で上書きする。
# ⚠️ 語は「文脈ごと」で書くこと。素の「操作」だけだと
# 『次世代デジタルサイネージ。スマホで簡単操作。』が手順書の扱いになった。
TITLE_PURPOSE = [
    (r"新卒|中途|募集|求人|インターン|バイトル|職種|エントリー", "採用（応募者の集客・入社前の理解）"),
    (r"採用ピッチ|会社説明会?", "採用（応募者の集客・入社前の理解）"),
    (r"導入事例|お客様の声|ユーザー(事例|の声)|活用事例",
     "サービス・商品の理解促進（検討中の見込み客向け）"),
    (r"操作(方法|手順|説明|マニュアル)|使い方|手順|チュートリアル|"
     r"How\s*to|ハウツー|設定方法|組立|取扱説明",
     "操作・業務手順の教育（問い合わせ削減）"),
    (r"\d+\s*周年|記念|ビジョン|理念|ブランドムービー|想い|パーパス|MVV",
     "ブランド・理念の発信（信頼の醸成）"),
    (r"展示会|ブース|出展", "展示会ブースの集客（通行者の足止め）"),
    (r"社内報", "社内への情報共有（理念・活動の浸透）"),
    (r"研修|教育|トレーニング|eラーニング", "社内教育（研修の標準化）"),
    (r"Official\s*(Music\s*)?Video|\bMV\b|M/V|Lyric\s*Video|Teaser|ティザー|"
     r"ミュージックビデオ", "楽曲のプロモーション（音楽作品の発表）"),
    (r"TVCM|WebCM|WEBCM|\bCM\b|スポット", "商品・サービスの認知獲得（広告出稿）"),
    (r"サービス紹介|商品紹介|製品紹介|機能紹介", "サービス・商品の理解促進（検討中の見込み客向け）"),
    (r"工場|技術|設備|ものづくり|製造|生産ライン",
     "技術力・設備の提示（取引先と採用の両方に効く）"),
    (r"会社案内|会社紹介|企業紹介", "会社の信頼獲得（営業・採用の共通素材）"),
]

CORP = r"(?:株式会社|\（株\）|\(株\)|有限会社|合同会社|一般社団法人|学校法人|医療法人)"
SEP = r"[^\s　｜|/／（）\(\)【】\[\]]"
CLIENT_PATS = [
    re.compile(rf"({CORP}{SEP}{{1,20}})\s*様"),
    re.compile(rf"({SEP}{{1,20}}{CORP})\s*様"),
    re.compile(rf"({SEP}{{2,20}})\s*(?:様|さま)"),
    re.compile(rf"({CORP}{SEP}{{1,20}})"),
    re.compile(rf"({SEP}{{1,20}}{CORP})"),
]
# 投稿者が制作会社か（＝発注企業ではない）。
VENDOR_RE = re.compile(
    r"動画制作|映像制作|映像製作|動画編集|制作会社|制作実績|制作事例|"
    r"クリエイティ|CREATIVE|Creative|creative|プロダクション|PRODUCTION|"
    r"Production|Films?\b|FILM|フィルム|スタジオ|STUDIO|Studio|"
    r"ビデオブレイン|Video\s*BRAIN|プルークス|PROOX|LOCUS|ロカス|"
    r"動画マーケティング|広告代理|デザイン事務所|ムービー制作|"
    r"Videoproduktion|Filmproduktion|映像屋|動画屋|PR動画|"
    r"動画製作所|映像製作所|動画工房|映像工房|シネマ|Cinema|cinema|"
    r"CHANGE|Analysis|VideoEdit|voss\s*create", re.I)
NOISE = re.compile(r"^(この|その|あの|弊社|当社|御社|貴社|皆|みな|本|同|今回|"
                   r"以下|上記|新|旧|全|各)")
# 題名に書かれた制作会社（「（LOCUS制作実績）」「【Video BRAIN制作】」）
VENDOR_IN_TITLE = re.compile(
    rf"[（(【]\s*({SEP}{{2,24}}?)\s*(?:制作実績|制作事例|制作|撮影編集)\s*[）)】]")
# 「（企業PR動画制作）」の『企業PR動画』のような、会社名でない説明句を弾く。
# 社名らしさ＝株式会社等が付く／英字の固有名／カタカナの屋号。
GENERIC_VENDOR_RE = re.compile(
    r"^(?:企業|会社|法人|自社|弊社|各種|他)?"
    r"(?:PR|CM|VP|AI|SNS|WEB|Web|採用|紹介|プロモーション|イメージ|広報|"
    r"インタビュー|アニメ|アニメーション|実写|縦型|短尺|映像|動画)*"
    r"(?:動画|映像|ムービー|CM|PV|MV)?$")
# 題名が「制作事例【靴屋Orike】…」の形のとき、【】の中が発注側。
BRACKET_CLIENT_RE = re.compile(
    r"(?:制作事例|制作実績|事例紹介|導入事例|実績紹介)\s*[【\[]\s*"
    r"([^】\]]{2,24})\s*[】\]]")


def norm_genre(topic: str, title: str = "") -> str:
    # お題を複数まとめて頼まれた日のレポートは、1行ごとの行き先が
    # お題からは決まらない（表の並び順で先頭の種類に全部寄ってしまい、
    # 会社紹介の事例が「MV」になった）。その時だけ題名で決める。
    if title and re.search(r"[、，]", topic):
        for name, pat in GENRE_RULES:
            if name != MISC and re.search(pat, title, re.I):
                return name
    for name, pat in GENRE_RULES:
        if re.search(pat, topic, re.I):
            return name
    return f"その他（{topic}）"


def guess_purpose(genre: str, title: str):
    for pat, purpose in TITLE_PURPOSE:
        m = re.search(pat, title, re.I)
        if m:
            return purpose, m.group(0)
    return GENRE_PURPOSE.get(genre, "不明"), "（種類からの既定）"


def guess_client(title: str, channel: str) -> str:
    """発注側の企業名。取れなければ空。制作会社の名前は返さない
    （制作会社を発注企業として出すと、提案先を間違える）。"""
    bm = BRACKET_CLIENT_RE.search(title)
    if bm:
        name = bm.group(1).strip("　 ・|｜_＿-")
        if (len(name) >= 2 and not NOISE.match(name)
                and not VENDOR_RE.search(name)
                and not GENERIC_VENDOR_RE.match(name)):
            return name
    for pat in CLIENT_PATS:
        for m in pat.finditer(title):
            name = m.group(1).strip("　 ・|｜_＿-")
            if len(name) < 2 or NOISE.match(name) or VENDOR_RE.search(name):
                continue
            return name
    if not VENDOR_RE.search(channel):
        return channel          # 企業・アーティストが自分で上げている
    return ""


def guess_vendor(title: str, channel: str) -> str:
    m = VENDOR_IN_TITLE.search(title)
    if m:
        name = m.group(1).strip("　 ・|｜_＿-")
        # 会社名でない説明句（「企業PR動画」「AI動画」）は制作会社ではない
        if (len(name) >= 2 and not NOISE.match(name)
                and not GENERIC_VENDOR_RE.match(name)):
            return name
    return channel if VENDOR_RE.search(channel) else ""


# 投稿者が個人の営業チャンネル。制作事例の棚には制作会社が普通に居るので、
# 「制作会社だから外す」ではなく「営業・売り込みそのもの」だけを外す。
SELF_PROMO_CH_RE = re.compile(
    r"ポートフォリオ|portfolio|フリーランス|案件(募集|承り)|お仕事(募集|承り)|"
    r"駆け出し|副業", re.I)


def out_of_scope(genre: str, title: str, dur, verdict: str, channel=""):
    """提案資料の材料にしない理由。無ければ None。"""
    if genre == MISC:
        return "企業映像でない（急上昇・題材の取り違え）"
    if verdict == "違う":
        return "視聴して別物と判定済み"
    if bot._VOICE_SAMPLE_RE.search(title) or bot._VOICE_SAMPLE_RE.search(channel):
        return "ボイスサンプル・ナレーターの営業"
    if SELF_PROMO_CH_RE.search(channel):
        return "制作者個人の営業チャンネル"
    for label, pat in bot._NOT_PROMO_PATTERNS:
        if pat.search(title):
            return label
    if dur and int(dur) > bot._LEAD_GEN_MIN_SEC and bot._LEAD_GEN_RE.search(title):
        return "集客のノウハウ解説"
    return None


def collect():
    rows, unreadable = [], Counter()
    for f in sorted(INSIGHTS.glob("*.md")):
        text = f.read_text(encoding="utf-8", errors="replace")
        mt = TOPIC_RE.search(text)
        topic = (mt.group("t").strip() if mt else "") or "急上昇"
        md = DATE_RE.search(text) or DATE_RE.search(f.name)
        date = md.group(1) if md else ""
        heads = [(m.start(), m.group(0))
                 for m in re.finditer(r"^## .+$", text, re.M)]

        for um in URL_RE.finditer(text):
            head = ""
            for pos, h in heads:
                if pos < um.start():
                    head = h
                else:
                    break
            hm = HEAD_RE.match(head)
            if not hm:
                unreadable[head[:50]] += 1
                continue
            title, channel = hm.group("title").strip(), hm.group("ch").strip()
            genre = norm_genre(topic, title)
            vm = META_RE.search(text[max(0, um.start() - 200):um.start() + 300])
            verdict = (vm.group("verdict") if vm else "") or "未判定"
            dur = hm.group("dur") or (vm.group("dur") if vm else "") or ""
            purpose, basis = guess_purpose(genre, title)
            rows.append({
                "動画の種類": genre,
                "発注企業・出演企業（推定）": guess_client(title, channel),
                "動画の目的（推定）": purpose,
                "推定の根拠": basis,
                "制作会社（推定）": guess_vendor(title, channel),
                "投稿チャンネル": channel,
                "題名": title,
                "尺（秒）": int(dur) if dur else "",
                "再生数": int(hm.group("views").replace(",", "")),
                "視聴判定": verdict,
                "リサーチ日": date,
                "検索お題": topic,
                "参考外の理由": out_of_scope(genre, title, dur, verdict, channel) or "",
                "URL": f"https://www.youtube.com/watch?v={um.group(1)}",
            })

    # 同じ動画が巡回で再登場する。URLで1本に寄せ、判定が付いている方を残す。
    uniq, seen = [], {}
    for r in rows:
        k = r["URL"]
        if k in seen:
            if uniq[seen[k]]["視聴判定"] == "未判定" != r["視聴判定"]:
                uniq[seen[k]] = r
            continue
        seen[k] = len(uniq)
        uniq.append(r)
    return uniq, len(rows), unreadable


def write_xlsx(rows, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    use = [r for r in rows if not r["参考外の理由"]]
    out = [r for r in rows if r["参考外の理由"]]
    cols = list(rows[0].keys())
    head_fill = PatternFill("solid", fgColor="1F3864")
    est_fill = PatternFill("solid", fgColor="7F7F7F")
    head_font = Font(color="FFFFFF", bold=True)

    wb = Workbook()

    def sheet(ws, data, title):
        ws.title = title
        ws.append(cols)
        for i, c in enumerate(cols, 1):
            cell = ws.cell(row=1, column=i)
            cell.fill = est_fill if "推定" in c else head_fill
            cell.font = head_font
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        for r in data:
            ws.append([r[c] for c in cols])
        widths = {"動画の種類": 24, "発注企業・出演企業（推定）": 26,
                  "動画の目的（推定）": 34, "推定の根拠": 16, "制作会社（推定）": 24,
                  "投稿チャンネル": 30, "題名": 60, "尺（秒）": 9, "再生数": 12,
                  "視聴判定": 10, "リサーチ日": 12, "検索お題": 22,
                  "参考外の理由": 26, "URL": 44}
        for i, c in enumerate(cols, 1):
            ws.column_dimensions[get_column_letter(i)].width = widths.get(c, 18)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(cols))}{len(data)+1}"
        for row in ws.iter_rows(min_row=2, max_row=len(data) + 1):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=False)
        return ws

    sheet(wb.active, use, "一覧")

    # 種類別のまとめ
    ws = wb.create_sheet("種類別まとめ")
    ws.append(["動画の種類", "本数", "再生数の中位", "尺の中位（秒）",
               "尺が分かる本数", "主な目的（推定・最多）", "制作会社が投稿した割合"])
    for i in range(1, 8):
        ws.cell(row=1, column=i).fill = head_fill
        ws.cell(row=1, column=i).font = head_font
    by = defaultdict(list)
    for r in use:
        by[r["動画の種類"]].append(r)
    for g, rs in sorted(by.items(), key=lambda kv: -len(kv[1])):
        durs = [r["尺（秒）"] for r in rs if r["尺（秒）"] != ""]
        pur = Counter(r["動画の目的（推定）"] for r in rs).most_common(1)[0]
        vend = sum(1 for r in rs if r["制作会社（推定）"])
        ws.append([g, len(rs),
                   int(statistics.median([r["再生数"] for r in rs])),
                   int(statistics.median(durs)) if durs else "",
                   len(durs), f"{pur[0]}（{pur[1]}本）",
                   f"{vend*100//len(rs)}%"])
    for i, w in enumerate([26, 8, 14, 15, 15, 46, 20], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"

    sheet(wb.create_sheet(), out, "参考外")

    # 読み方
    ws = wb.create_sheet("この表の読み方")
    ws.column_dimensions["A"].width = 110
    for line in [
        "■ この表は何か",
        f"　Discordボットが {rows[0]['リサーチ日']} 以降に毎日集めた YouTube リサーチ"
        f"（レポート {len(list(INSIGHTS.glob('*.md')))} 本）から、動画を1本1行に並べ直したもの。",
        "",
        "■ 事実と推定を分けてある（灰色の見出しが推定）",
        "　【事実】題名・投稿チャンネル・再生数・尺・URL・視聴判定・リサーチ日・検索お題",
        "　　　　　…YouTube API から取った値と、Gemini が実際に視聴して付けた判定。",
        "　【推定】発注企業・出演企業／動画の目的／制作会社",
        "　　　　　…レポート本文に「目的」は書かれていない。題名の語と検索お題から",
        "　　　　　　機械的に導いている。当たった語は『推定の根拠』列にある。",
        "",
        "　⚠️ 推定列をそのまま「調査結果」としてお客様に出さないこと。",
        "　　 提案に使う分だけ、URL を開いて中身を見てから書くこと。",
        "",
        "■ 視聴判定（10月以降のレポートのみ付いている）",
        "　実物　…検索お題どおりの動画だった",
        "　近い　…隣接するジャンルだった",
        "　違う　…別物だった → 『参考外』シートへ送っている",
        "　未判定…判定の仕組みを入れる前（9月まで）のレポート",
        "",
        "■ 参考外シートに送った理由",
        "　ボイスサンプル／BGM見本／制作者のポートフォリオ／作り方の解説／",
        "　ショート動画／ゲーム実況／急上昇の雑多な動画／視聴して別物と判定したもの。",
        "　判定の規則は discord-groupchat/ai_group_chat.py のものをそのまま使っている。",
        "",
        "■ 作り直し方",
        "　cd ~/so-portfolio/discord-groupchat",
        "　./venv/bin/python tools/build_proposal_table.py",
    ]:
        ws.append([line])
        c = ws.cell(row=ws.max_row, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if line.startswith("■"):
            c.font = Font(bold=True, color="1F3864")
        if line.strip().startswith("⚠️"):
            c.font = Font(bold=True, color="C00000")

    wb.save(path)
    return len(use), len(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--dry", action="store_true")
    args = ap.parse_args()

    rows, raw, unreadable = collect()
    n_rep = len(list(INSIGHTS.glob("*.md")))
    print(f"レポート {n_rep} 本 → 動画 {raw} 件 → 重複を除いて {len(rows)} 件")
    if unreadable:
        print("⚠️ 見出しが読めなかった:")
        for k, v in unreadable.most_common(10):
            print(f"   {v}件 {k}")

    use = [r for r in rows if not r["参考外の理由"]]
    print(f"提案に使える {len(use)} 件 / 参考外 {len(rows)-len(use)} 件")
    print("\n参考外の理由:")
    for k, v in Counter(r["参考外の理由"] for r in rows
                        if r["参考外の理由"]).most_common():
        print(f"  {v:4d}  {k}")
    print("\n種類ごとの本数（提案に使える分）:")
    for k, v in Counter(r["動画の種類"] for r in use).most_common():
        print(f"  {v:4d}  {k}")
    named = sum(1 for r in use if r["発注企業・出演企業（推定）"])
    vend = sum(1 for r in use if r["制作会社（推定）"])
    basis = Counter(r["推定の根拠"] for r in use)
    print(f"\n企業名が取れた: {named}/{len(use)} ({named*100//len(use)}%)")
    print(f"制作会社が取れた: {vend}/{len(use)} ({vend*100//len(use)}%)")
    print(f"目的が題名の語から取れた: {len(use)-basis['（種類からの既定）']}/{len(use)}"
          f" （残りは種類からの既定）")
    if args.dry:
        return 0

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    cols = list(rows[0].keys())
    csv_path = OUT_DIR / "リサーチ一覧.csv"
    with csv_path.open("w", encoding="utf-8-sig", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=cols)
        w.writeheader()
        w.writerows(rows)
    xlsx_path = OUT_DIR / "YouTubeリサーチ一覧.xlsx"
    n_use, n_out = write_xlsx(rows, xlsx_path)
    print(f"\n書き出し:\n  {xlsx_path}（一覧 {n_use} / 参考外 {n_out}）\n  {csv_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
